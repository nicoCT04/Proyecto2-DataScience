"""Arma el informe final a partir de las secciones en Markdown, con formato APA 7.

Uso desde la raíz:  python informe/build_informe.py [--pdf]

Lee informe/secciones/*.md en orden, renumera los títulos de primer nivel (1, 2, 3...)
y genera Informe_Proyecto2_Resultados.docx. Con --pdf lo exporta también a PDF mediante
Microsoft Word (docx2pdf).

Formato APA 7: Times New Roman 12 en negro, doble espacio, márgenes de 2.54 cm, sangría
de 1.27 cm en la primera línea de cada párrafo, títulos de nivel 1 centrados en negrita,
nivel 2 alineados a la izquierda en negrita y nivel 3 en negrita cursiva; tablas y figuras
con su número en negrita y el título en cursiva encima; tablas solo con líneas
horizontales; número de página arriba a la derecha; referencias en página nueva con
sangría francesa.

Markdown soportado: # / ## / ### títulos, párrafos, listas con "- " y "1. ", tablas con
"|", imágenes ![Figura N. Título](ruta relativa al .md), **negrita**, *cursiva* y
`código`. Un párrafo en cursiva que empieza con "Tabla N." es el título de la tabla
siguiente.
"""
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SECTIONS = ROOT / "informe" / "secciones"
OUT = ROOT / "Informe_Proyecto2_Resultados.docx"
LOGO = ROOT / "informe" / "assets" / "uvg_logo.jpg"

INTEGRANTES = [("Carlos Nicolás Concuá Trujillo", "23197"), ("Jorge Luis Felipe Aguilar Portillo", "23195"),
               ("Fernando José Rueda Rodas", "23748"), ("Fernando Andree Hernández Martínez", "23645")]
REPO = "https://github.com/nicoCT04/Proyecto2-DataScience"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre"]
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
INDENT = Cm(1.27)
TEXT_WIDTH = Cm(16.51)   # carta menos 2 × 2.54 cm
MAX_FIG_HEIGHT = Cm(17)  # deja lugar al número y título de la figura en la misma página


# ----------------------------------------------------------------------------
# Estilos
# ----------------------------------------------------------------------------
def set_font(obj, size=12, bold=None, italic=None):
    """Times New Roman negro en un estilo o run, también para el texto de Asia/complejo."""
    f = obj.font
    f.name, f.size, f.color.rgb = FONT, Pt(size), BLACK
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    rpr = obj.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(theme), None)  # los temas de Word pisarían la fuente


def double(pf, indent=True):
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = pf.space_after = Pt(0)
    pf.first_line_indent = INDENT if indent else Cm(0)


def setup(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(27.94), Cm(21.59)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    normal = doc.styles["Normal"]
    set_font(normal)
    double(normal.paragraph_format)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name, align, italic in [("Heading 1", WD_ALIGN_PARAGRAPH.CENTER, False),
                                ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT, False),
                                ("Heading 3", WD_ALIGN_PARAGRAPH.LEFT, True)]:
        st = doc.styles[name]
        set_font(st, bold=True, italic=italic)
        double(st.paragraph_format, indent=False)
        st.paragraph_format.alignment = align
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        set_font(doc.styles[name], bold=False, italic=False)
        double(doc.styles[name].paragraph_format, indent=False)
    page_number_header(sec)


def page_number_header(section):
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_font(run)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------------------
# Portada (portada de estudiante APA 7 con el logo de la UVG)
# ----------------------------------------------------------------------------
def cover(doc):
    today = date.today()

    def line(text="", bold=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        if text:
            set_font(p.add_run(text), bold=bold)
        return p

    line().add_run().add_picture(str(LOGO), width=Cm(4.5))
    line()
    line("Proyecto 2 — Resultados: Clasificación de Acciones en Video de la Bundesliga", bold=True)
    line("(Pase, Disputa, Saque de Banda y Fondo)", bold=True)
    line()
    for name, carne in INTEGRANTES:
        line(f"{name} ({carne})")
    line("Departamento de Ciencias de la Computación, Universidad del Valle de Guatemala")
    line("CC3084: Data Science")
    line(f"{today.day} de {MESES[today.month - 1]} de {today.year}")
    page_break(doc)


# ----------------------------------------------------------------------------
# Markdown -> docx
# ----------------------------------------------------------------------------
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(p, text, size=12, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            set_font(p.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("`"):
            set_font(p.add_run(part[1:-1]), size, bold=bold or None)
        elif part.startswith("*"):
            set_font(p.add_run(part[1:-1]), size, bold=bold or None, italic=True)
        else:
            set_font(p.add_run(part), size, bold=bold or None)


def label(doc, number, title):
    """Número en negrita y título en cursiva, encima de la tabla o figura (APA)."""
    for text, fmt in [(number, "**"), (title, "*")]:
        if text:
            p = doc.add_paragraph()
            double(p.paragraph_format, indent=False)
            p.paragraph_format.keep_with_next = True
            add_inline(p, f"{fmt}{text}{fmt}")


def split_caption(caption):
    m = re.match(r"^((?:Tabla|Figura) \d+)\.\s*(.*)$", caption.strip())
    return (m.group(1), m.group(2)) if m else ("", caption)


def border(cell, edges):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "8")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcpr.append(borders)


def add_table(doc, rows, caption):
    if caption:
        label(doc, *split_caption(caption))
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], cells[2:]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 11 if len(header) <= 4 else (10 if len(header) <= 6 else 9)
    for i, row in enumerate([header] + body):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
            add_inline(p, text, size=size, bold=(i == 0))
            border(cell, {"top", "bottom"} if i == 0 else ({"bottom"} if i == len(body) else set()))
    double(doc.add_paragraph().paragraph_format, indent=False)


def add_image(doc, md_path, caption, target):
    img = (md_path.parent / target).resolve()
    label(doc, *split_caption(caption))
    w, h = Image.open(img).size
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(12)
    p.add_run().add_picture(str(img), width=min(TEXT_WIDTH, int(MAX_FIG_HEIGHT * w / h)))


def render(doc, md_path, number):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, para, table_caption = 0, [], None

    def flush():
        if para:
            add_inline(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            flush()
        elif ln.startswith("#"):
            flush()
            level = len(ln) - len(ln.lstrip("#"))
            title = re.sub(r"^\d+(\.\d+)*\.?\s+", "", ln.lstrip("#").strip())
            num = re.match(r"^#+\s+\d+\.(\d+(\.\d+)*)", ln)
            prefix = f"{number}." + num.group(1) if (num and level > 1) else (f"{number}." if level == 1 else "")
            doc.add_heading(f"{prefix} {title}".strip(), level=level)
        elif ln.startswith("!["):
            flush()
            m = re.match(r"!\[(.*)\]\((.*)\)", ln)
            add_image(doc, md_path, m.group(1), m.group(2))
        elif ln.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows, table_caption)
            table_caption = None
            continue
        elif re.match(r"^\s*(- |\d+\. )", ln):
            flush()
            numbered = bool(re.match(r"^\s*\d+\. ", ln))
            text = re.sub(r"^\s*(- |\d+\. )", "", ln)
            i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
                text += " " + lines[i].strip()
                i += 1
            add_inline(doc.add_paragraph(style="List Number" if numbered else "List Bullet"), text)
            continue
        elif re.fullmatch(r"\*Tabla [^*]+\*", ln.strip()):
            flush()
            table_caption = ln.strip()[1:-1]
        else:
            para.append(ln.strip())
        i += 1
    flush()


def references(doc, md_path, number):
    """Referencias en página nueva, con sangría francesa y doble espacio (APA)."""
    page_break(doc)
    doc.add_heading(f"{number}. Referencias", level=1)
    text = md_path.read_text(encoding="utf-8")
    for entry in [e.strip() for e in text.split("\n\n")[1:] if e.strip()]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = INDENT
        p.paragraph_format.first_line_indent = -INDENT
        add_inline(p, entry.replace("\n", " "))


def main():
    doc = Document()
    setup(doc)
    cover(doc)
    add_inline(doc.add_paragraph(), f"Repositorio: {REPO}. Aplicación: `streamlit run app/app.py`.")
    files = sorted(SECTIONS.glob("*.md"))
    for n, f in enumerate(files, 1):
        if "referencias" in f.name:
            references(doc, f, n)
        else:
            render(doc, f, n)
    doc.save(OUT)
    print(f"Guardado {OUT.relative_to(ROOT)} ({len(files)} secciones)")
    if "--pdf" in sys.argv:
        from docx2pdf import convert
        convert(str(OUT), str(OUT.with_suffix(".pdf")))
        print(f"Guardado {OUT.with_suffix('.pdf').relative_to(ROOT)}")


if __name__ == "__main__":
    main()
